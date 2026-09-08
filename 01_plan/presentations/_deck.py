"""Layout grammar for OFF-ROAD EV PLATFORM decks, on top of python-pptx.

Implements the DYN-09 structure — photo strip, blue title band, body, blue
footer, chevron badge in the corner — as five reusable slide layouts, so the
content file only ever says what goes on a slide, never where the band is.

    std(title, strip)   photo strip + title band + white body   (DYN-09 exact)
    tall(title)         title band at the top, full-height body
    split(title, img)   title band + a full-bleed image column
    cover(...)          full-bleed hero + solid block
    dark(title)         full deep-blue statement slide
"""
from __future__ import annotations

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _brand as B

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
ASSETS = os.path.join(HERE, "assets")
FIG = os.path.join(ASSETS, "figures")
CROP = os.path.join(ASSETS, "crops")
WEB = os.path.join(ASSETS, "web")
SHOT = os.path.join(ROOT, "03_hardware", "vesc", "screenshots")

C = lambda h: RGBColor(*B.rgb(h))          # noqa: E731


class Deck:
    def __init__(self, footer_left):
        self.prs = Presentation()
        self.prs.slide_width = Inches(B.W)
        self.prs.slide_height = Inches(B.H)
        self.footer_left = footer_left
        self.n = 0

    # -- primitives --------------------------------------------------------
    def _blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def rect(self, sl, x, y, w, h, fill, line=None, lw=0.75,
             shape=MSO_SHAPE.RECTANGLE):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = C(fill)
        if line:
            s.line.color.rgb = C(line)
            s.line.width = Pt(lw)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        s.text_frame.word_wrap = True
        return s

    def img(self, sl, path, x, y, w, h):
        """object-fit: cover — fill the rect, crop the overflow, never stretch."""
        iw, ih = Image.open(path).size
        na, ta = iw / ih, w / h
        pic = sl.shapes.add_picture(path, Inches(x), Inches(y),
                                    Inches(w), Inches(h))
        if na > ta:
            f = (1 - ta / na) / 2
            pic.crop_left = pic.crop_right = f
        elif na < ta:
            f = (1 - na / ta) / 2
            pic.crop_top = pic.crop_bottom = f
        return pic

    def fit(self, sl, path, x, y, w, h):
        """object-fit: contain — whole image inside the rect, centred."""
        iw, ih = Image.open(path).size
        na, ta = iw / ih, w / h
        if na > ta:
            ww, hh = w, w / na
        else:
            hh, ww = h, h * na
        return sl.shapes.add_picture(path, Inches(x + (w - ww) / 2),
                                     Inches(y + (h - hh) / 2),
                                     Inches(ww), Inches(hh))

    def text(self, sl, x, y, w, h, paras, size=None, color=B.INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, line=1.15,
             font=None, italic=False):
        """paras: str, or list of str, or list of dicts {t,size,color,bold,...}."""
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if isinstance(paras, str):
            paras = [paras]
        for i, item in enumerate(paras):
            d = {"t": item} if isinstance(item, str) else dict(item)
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = d.get("align", align)
            p.line_spacing = d.get("line", line)
            if i:
                p.space_before = Pt(d.get("space", space))
            r = p.add_run()
            r.text = d["t"]
            f = r.font
            f.name = d.get("font", font or B.FONT)
            f.size = Pt(d.get("size", size or B.SZ_BODY))
            f.bold = d.get("bold", bold)
            f.italic = d.get("italic", italic)
            f.color.rgb = C(d.get("color", color))
        return tb

    def chip(self, sl, x, y, label, color, w=None):
        w = w or (0.13 + 0.062 * len(label))
        s = self.rect(sl, x, y, w, 0.17, color,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        s.adjustments[0] = 0.34
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = B.FONT
        r.font.size = Pt(B.SZ_TAG)
        r.font.bold = True
        r.font.color.rgb = C(B.WHITE)
        return s

    def card(self, sl, x, y, w, h, head, body, fill=B.PANEL, accent=B.BLUE,
             hs=B.SZ_H2, bs=B.SZ_BODY, pad=0.20, tc=None):
        self.rect(sl, x, y, w, h, fill, line=B.LINE if fill == B.WHITE else None)
        self.text(sl, x + pad, y + pad - 0.03, w - 2 * pad, 0.3, head,
                  size=hs, color=accent, bold=True)
        if body:
            self.text(sl, x + pad, y + pad + 0.26, w - 2 * pad,
                      h - 2 * pad - 0.26, body, size=bs,
                      color=tc or B.INK, line=1.22, space=5)

    def stat(self, sl, x, y, w, value, label, color=B.BLUE, vs=B.SZ_STAT):
        self.text(sl, x, y, w, 0.48, value, size=vs, color=color, bold=True)
        self.text(sl, x, y + 0.44, w, 0.34, label, size=B.SZ_SMALL,
                  color=B.GREY, line=1.15)

    # -- chrome ------------------------------------------------------------
    def _badge(self, sl, w=B.BADGE_W, h=B.BADGE_H):
        sl.shapes.add_picture(os.path.join(FIG, "badge.png"), Inches(0),
                              Inches(0), Inches(w), Inches(h))

    def _footer(self, sl, dark=False):
        if not dark:
            self.rect(sl, 0, B.FOOT_Y, B.W, B.FOOT_H, B.BLUE)
        self.text(sl, B.M, B.FOOT_Y + 0.16, 6.4, 0.22, self.footer_left,
                  size=B.SZ_FOOT, color=B.WHITE)
        self.text(sl, B.W - B.M - 1.0, B.FOOT_Y + 0.16, 1.0, 0.22,
                  f"{self.n:02d}", size=B.SZ_FOOT, color=B.WHITE,
                  align=PP_ALIGN.RIGHT)

    def _band(self, sl, y, title, x_text):
        self.rect(sl, 0, y, B.W, B.BAND_H, B.BLUE)
        # The band is one line tall by design. Long titles step down rather
        # than wrap and get clipped by it.
        size = B.SZ_TITLE if len(title) <= 42 else (23 if len(title) <= 52 else 20)
        self.text(sl, x_text, y, B.W - x_text - B.M, B.BAND_H, title,
                  size=size, color=B.WHITE, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE)

    # -- layouts -----------------------------------------------------------
    def std(self, title, strip):
        """DYN-09 exact: photo strip, title band, white body. Body is short."""
        self.n += 1
        sl = self._blank()
        self.img(sl, os.path.join(CROP, strip), 0, 0, B.W, B.STRIP_H)
        self._band(sl, B.BAND_Y, title, B.M)
        self._badge(sl)
        self._footer(sl)
        return sl, (B.M, B.BODY_Y + 0.22, B.CW, B.BODY_H - 0.34)

    def tall(self, title):
        """Title band at the very top. Full-height body for figures and tables."""
        self.n += 1
        sl = self._blank()
        self._band(sl, 0, title, 1.02)
        self._badge(sl, 0.78, 0.88)
        self._footer(sl)
        return sl, (B.M, 1.02, B.CW, B.FOOT_Y - 1.02 - 0.14)

    def split(self, title, image, side="left", frac=0.36):
        """Title band + a full-bleed image column. Returns the content rect."""
        self.n += 1
        sl = self._blank()
        iw = B.W * frac
        top = B.BAND_H
        ix = 0 if side == "left" else B.W - iw
        self.img(sl, image, ix, top, iw, B.FOOT_Y - top)
        self._band(sl, 0, title, 1.02)
        self._badge(sl, 0.78, 0.88)
        self._footer(sl)
        if side == "left":
            return sl, (iw + 0.34, top + 0.30, B.W - iw - 0.34 - B.M,
                        B.FOOT_Y - top - 0.46)
        return sl, (B.M, top + 0.30, B.W - iw - 0.34 - B.M,
                    B.FOOT_Y - top - 0.46)

    def cover(self, hero, title, subtitle, meta):
        """Photo left, solid block right — the DYN-09 cover, and the shape that
        suits a portrait photograph without cropping the vehicle in half."""
        self.n += 1
        sl = self._blank()
        px = 4.55
        self.img(sl, hero, 0, 0, px, B.H)
        self.rect(sl, px, 0, B.W - px, B.H, B.BLUE_DK)
        self._badge(sl)
        tx, tw = px + 0.58, B.W - px - 0.58 - 0.50
        self.text(sl, tx, 1.42, tw, 1.30, title, size=33, color=B.WHITE,
                  bold=True, line=1.06)
        self.text(sl, tx, 2.92, tw, 0.72, subtitle, size=12.5, color=B.WHITE,
                  line=1.30)
        self.text(sl, tx, 4.02, tw, 0.60, meta, size=10, color=B.CYAN,
                  bold=True, line=1.45)
        return sl

    def dark(self, title):
        self.n += 1
        sl = self._blank()
        self.rect(sl, 0, 0, B.W, B.H, B.BLUE_DK)
        self._band(sl, 0, title, 1.02)
        self._badge(sl, 0.78, 0.88)
        self._footer(sl, dark=True)
        return sl, (B.M, 1.10, B.CW, B.FOOT_Y - 1.10 - 0.14)

    def save(self, path):
        self.prs.save(path)
        return path
