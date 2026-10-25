"""Number audit: every number on a slide must resolve to the repo.

The review-deck skill makes this mandatory before presenting: "a number that
appears only in the deck is a fabricated number until proven otherwise". This
does the grep for you. For each number on each slide it reports the first
place in the sources of truth where a matching value lives, or NOT FOUND.

Matching is by value at the slide's precision (34.2 on a slide matches 34.16
in a results file), so read the reported location: a match is a lead, not a
proof. NOT FOUND is either a fabricated number or one the deck code derives
from the params - check which, and say so in the deck code.

    ./.venv/Scripts/python.exe 01_plan/presentations/_audit.py 01_plan/presentations/C01-review.pptx
"""
from __future__ import annotations

import glob
import os
import re
import sys

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))

# Sources of truth (skill: 03_hardware, 04_model/params, 06_data, the VESC XML)
# plus the model outputs. Curve CSVs are left out: thousands of samples would
# match any number at all.
SOURCES = [
    "03_hardware/*.md", "03_hardware/vesc/*.xml", "03_hardware/vesc/*.md",
    "04_model/params/*.m", "04_model/results/*.md",
    "04_model/results/C01_oat.csv", "04_model/results/C01_coast_start.csv",
    "06_data/**/*.md", "07_experiments/*.md",
]
NUM = re.compile(r"(?<![\w.])(\d{1,3}(?:[  ]\d{3})+|\d+(?:[.,]\d+)?)(?![\w.])")
TRIVIAL = set(range(0, 11)) | {2026, 2027}


def numbers(text):
    for m in NUM.finditer(text):
        raw = m.group(1)
        s = raw.replace(" ", "").replace(" ", "").replace(",", ".")
        dec = len(s.split(".")[1]) if "." in s else 0
        yield raw, float(s), dec, m.start()


def corpus():
    out = []
    for pat in SOURCES:
        for path in sorted(glob.glob(os.path.join(ROOT, pat), recursive=True)):
            rel = os.path.relpath(path, ROOT).replace("\\", "/")
            try:
                lines = open(path, encoding="utf-8").read().splitlines()
            except UnicodeDecodeError:
                continue
            for i, ln in enumerate(lines, 1):
                for _, v, _, _ in numbers(ln):
                    out.append((v, rel, i, ln.strip()))
                # scientific notation in .m files: 16.855e-3, 138.8e-6
                for m in re.finditer(r"(\d+(?:\.\d+)?)e(-?\d+)", ln):
                    v = float(m.group(1)) * 10 ** int(m.group(2))
                    for scale in (1, 1e3, 1e6):
                        out.append((v * scale, rel, i, ln.strip()))
    return out


def find(value, dec, corp):
    half = 0.5 * 10 ** (-dec)
    for v, rel, i, ln in corp:
        if abs(v - value) <= half + 1e-12:
            return f"{rel}:{i}"
    return None


def audit(deck):
    corp = corpus()
    prs = Presentation(deck)
    rows, missing = [], 0
    for n, slide in enumerate(prs.slides, 1):
        seen = set()
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            text = shp.text_frame.text
            if text.startswith("OFF-ROAD EV PLATFORM") or re.fullmatch(r"\d\d", text.strip()):
                continue        # footer and page number
            for raw, v, dec, pos in numbers(text):
                if dec == 0 and v in TRIVIAL:
                    continue
                if (raw, n) in seen:
                    continue
                seen.add((raw, n))
                ctx = text[max(0, pos - 18):pos + len(raw) + 18].replace("\n", " ")
                hit = find(v, dec, corp)
                missing += hit is None
                rows.append((n, raw, ctx, hit or "NOT FOUND"))
    w = max(len(r[2]) for r in rows) if rows else 10
    for n, raw, ctx, hit in rows:
        flag = "!!" if hit == "NOT FOUND" else "  "
        print(f"{flag} {n:>2}  {raw:>9}  {ctx:<{w}}  {hit}")
    print(f"\n{len(rows)} numbers checked, {missing} not found in the sources of truth.")
    return missing


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    deck = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "C01-review.pptx")
    audit(deck)
